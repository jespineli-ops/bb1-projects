#!/usr/bin/env python3
"""
Generates the BB1 Technical Notes document for the Quorum Pre-Billing Report
customisation.

Requirements: python-docx (pip3 install --user python-docx)

Note: this machine has no cairo/Node available, so the ERD and Process Flow
are rendered as BB1-styled tables (relationship table / swimlane table)
instead of SVG->PNG diagrams - the same fallback used by
generate_tech_notes.py (Tenancy Schedule Report) and
generate_tech_notes_customer_statement.py (Customer Statement) in this same
folder.

Usage: python3 generate_tech_notes_prebilling.py
Output: ../BB1_TechNotes_Quorum_PreBilling.docx (relative to this file)
"""

import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ---------------------------------------------------------------------------
# Brand constants
# ---------------------------------------------------------------------------
BB1_BLUE = RGBColor(0x24, 0x78, 0xAA)
BB1_RED = RGBColor(0xC0, 0x28, 0x2E)
BB1_GREY = RGBColor(0x9B, 0x96, 0x92)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
CODE_GREY = RGBColor(0x44, 0x44, 0x44)

BB1_BLUE_HEX = "2478AA"
GREY_LIGHT_HEX = "F5F5F5"
WHITE_HEX = "FFFFFF"
BORDER_HEX = "CCCCCC"

FONT = "Roboto"
MONO_FONT = "Courier New"

CONTENT_WIDTH_CM = 17.0  # A4 (21cm) minus 2cm margins each side
ZWSP = "​"

HERE = os.path.dirname(os.path.abspath(__file__))
OUTPUT_PATH = os.path.normpath(os.path.join(HERE, "..", "BB1_TechNotes_Quorum_PreBilling.docx"))


# ---------------------------------------------------------------------------
# Low-level oxml helpers (python-docx has no first-class API for cell
# shading, custom borders, or forced column widths)
# ---------------------------------------------------------------------------

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_borders(cell, hex_color=BORDER_HEX, sz=4):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), hex_color)
        borders.append(el)
    tcPr.append(borders)


def set_cell_width(cell, width_cm):
    cell.width = Cm(width_cm)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = OxmlElement('w:tcW')
    tcW.set(qn('w:w'), str(int(width_cm * 567)))  # 1cm = 567 twips
    tcW.set(qn('w:type'), 'dxa')
    tcPr.append(tcW)


def set_table_fixed_layout(table, total_width_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(int(total_width_cm * 567)))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)


def set_table_grid(table, col_widths_cm):
    """python-docx's add_table() writes a tblGrid with the columns divided evenly,
    which does not automatically follow later per-cell set_cell_width() calls. Some
    renderers (Google Docs, LibreOffice) lay out columns from tblGrid rather than the
    individual w:tcW on each cell, which would silently undo the custom widths set
    below - so the grid itself must be rewritten to match."""
    tbl = table._tbl
    grid = tbl.find(qn('w:tblGrid'))
    for col in list(grid):
        grid.remove(col)
    for width_cm in col_widths_cm:
        gridCol = OxmlElement('w:gridCol')
        gridCol.set(qn('w:w'), str(int(width_cm * 567)))
        grid.append(gridCol)


def set_cell_margins(cell, top=40, bottom=40, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement('w:tcMar')
    for edge, val in (('top', top), ('bottom', bottom), ('left', left), ('right', right)):
        node = OxmlElement(f'w:{edge}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        mar.append(node)
    tcPr.append(mar)


def soft_break(text):
    """Zero-width space after every '_' and '.' in long tokens (>=15 chars),
    so Word can wrap long NetSuite internal ids inside a narrow table cell."""
    text = str(text)
    parts = text.split(' ')
    out = []
    for part in parts:
        if len(part) >= 15:
            part = part.replace('_', '_' + ZWSP).replace('.', '.' + ZWSP)
        out.append(part)
    return ' '.join(out)


# ---------------------------------------------------------------------------
# Content helpers
# ---------------------------------------------------------------------------

def set_default_font(doc):
    style = doc.styles['Normal']
    style.font.name = FONT
    style.font.size = Pt(10)
    rpr = style.element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)


def set_margins(doc):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)


def add_run(paragraph, text, size=10, bold=False, italic=False, color=BLACK, font=FONT):
    run = paragraph.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def add_heading(doc, text, level, number=None):
    """level: 1=H1, 2=H2 (red), 3=H3"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.outline_level = level - 1
    label = f"{number} {text}" if number else text
    size = {1: 16, 2: 14, 3: 13}[level]
    color = BB1_RED if level == 2 else BB1_BLUE
    add_run(p, label, size=size, bold=True, color=color)
    return p


def add_body(doc, text, size=10, italic=False, color=BLACK, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(space_after)
    add_run(p, text, size=size, italic=italic, color=color)
    return p


def add_hyperlink(paragraph, url, text, color=BB1_BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                           is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rPr.append(rFonts)
    color_el = OxmlElement('w:color')
    color_el.set(qn('w:val'), '2478AA')
    rPr.append(color_el)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '20')
    rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement('w:t')
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_table(doc, headers, rows, col_widths_cm, mono_cols=None, note=None):
    """headers: list[str]; rows: list[list[str]]; col_widths_cm sums to CONTENT_WIDTH_CM."""
    mono_cols = mono_cols or set()
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_fixed_layout(table, sum(col_widths_cm))
    set_table_grid(table, col_widths_cm)
    table.autofit = False

    # header row
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_width(hdr[i], col_widths_cm[i])
        set_cell_background(hdr[i], BB1_BLUE_HEX)
        set_cell_borders(hdr[i])
        set_cell_margins(hdr[i])
        hdr[i].paragraphs[0].text = ''
        add_run(hdr[i].paragraphs[0], h, size=11, bold=True, color=WHITE)

    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = WHITE_HEX if r_idx % 2 == 0 else GREY_LIGHT_HEX
        for c_idx, value in enumerate(row):
            cell = cells[c_idx]
            set_cell_width(cell, col_widths_cm[c_idx])
            set_cell_background(cell, fill)
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.paragraphs[0].text = ''
            values = value if isinstance(value, list) else [value]
            for v_idx, v in enumerate(values):
                p = cell.paragraphs[0] if v_idx == 0 else cell.add_paragraph()
                is_mono = c_idx in mono_cols
                # soft_break runs on every cell, not just monospace ones - long
                # NetSuite ids turn up inside ordinary prose columns too (e.g. a
                # "Detail" or "To" column), and an unbroken id there bleeds off
                # the page just as readily as one in a dedicated id column
                display = soft_break(v)
                add_run(p, display, size=9 if is_mono else 10,
                        color=CODE_GREY if is_mono else BLACK,
                        font=MONO_FONT if is_mono else FONT)

    if note:
        add_body(doc, note, size=8, italic=True, color=BB1_GREY, space_after=14)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(10)

    return table


def add_numbered_list(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Number')
        add_run(p, item, size=10)


# ---------------------------------------------------------------------------
# Build the document
# ---------------------------------------------------------------------------

def build():
    doc = Document()
    set_default_font(doc)
    set_margins(doc)

    # ---- Cover ----
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(120)
    add_run(p, "BB1", size=48, bold=True, color=BB1_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Technical Notes", size=28, bold=True, color=BB1_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    add_run(p, "Quorum", size=22, bold=True, color=BB1_RED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(p, "Pre-Billing Report (Tenant Billing History)", size=18, bold=True, color=BB1_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(40)
    add_run(p, "Author: Jared Espineli   |   Date: 17 September 2026", size=11, color=BB1_GREY)

    doc.add_page_break()

    # ---- 1. Overview ----
    add_heading(doc, "Overview", 1, "1.")

    add_heading(doc, "Introduction", 3, "1.1")
    add_body(doc,
             "The Pre-Billing Report gives Quorum's finance and billing team a way to check every tenant's "
             "billing history before the monthly billing run is executed, rather than discovering a mistake "
             "only after invoices have gone out. Modelled on the equivalent \"Pre-Billing Check\" report Quorum "
             "used in MRI, it is delivered as a single NetSuite Suitelet: a user narrows the run to a Property "
             "Portfolio, Property, Accommodation Type and/or Tenant, picks the billing period to review and how "
             "many months of history to include, and the report assembles each tenant's opening balance, "
             "itemised charges, receipts, credits and any bad debt movement into a running-balance ledger, with "
             "a portfolio-wide reconciliation check at the end. The result can be reviewed on screen, printed to "
             "a branded PDF, or exported to CSV, so any discrepancy - an invoice not agreeing with its posted "
             "amount, or a balance that doesn't prove - is caught and corrected before tenants are billed.")

    add_heading(doc, "BRS / Specification", 3, "1.2")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    add_run(p, "Service Request: ", size=10, bold=True)
    add_hyperlink(p, "https://docs.google.com/document/d/1qvcHPykKK6SChzqyXX9xkqxL_R0vBPxOOdIIwMmoFHA/edit?tab=t.0",
                  "Pre-Billing Report - Service Request")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    add_run(p2, "Related BRS (contract/charge data this report depends on): ", size=10, bold=True)
    add_hyperlink(p2, "https://docs.google.com/document/d/1SAVNacnxpZ0WIQhK6tY6jN8gzMs--dXqyMUz91YDz1s/edit?tab=t.0",
                  "Quorum Contracts")
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(4)
    add_run(p3, "Sample report referenced in the Service Request: ", size=10, bold=True)
    add_hyperlink(p3, "https://drive.google.com/file/d/1fRPM23PLNdYjXvQwfrdI-XzAG-JRSrDJ/view?usp=drive_link",
                  "Pre-Billing Report (Drive file)")
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(12)
    add_run(p4, "Teamwork task: ", size=10, bold=True)
    add_hyperlink(p4, "https://teamwork.bluebridgeone.com/app/tasks/43431016", "43431016")

    add_heading(doc, "Development Environment", 3, "1.3")
    add_body(doc, "Quorum Production - NetSuite Account 11536405. This customisation was built and is deployed "
                  "directly in the production account; there is no separate sandbox deployment for it.")

    add_heading(doc, "Business Requirements", 2, "1.4")

    add_heading(doc, "Summary", 3, "1.4.1")
    add_body(doc,
             "Before Quorum runs its monthly tenant billing, the finance team needs confidence that every "
             "tenant's account is accurate - that charges, receipts and credits already posted agree with what "
             "will appear on the next invoice, and that no balance has drifted out of reconciliation. The "
             "Pre-Billing Report customisation gives them a self-service way to check this directly in NetSuite: "
             "select the portfolio, property or tenant to review, choose how many billing periods of history to "
             "look back over, and the system produces a period-by-period ledger for every matching tenant, "
             "complete with a running balance and a portfolio-wide total that proves the opening balance plus "
             "all movement equals the closing balance. Anything that would otherwise only surface as a tenant "
             "dispute or a retroactive credit note after billing - an unreconciled invoice, an unexpected write-"
             "off, a balance that doesn't prove - is instead caught and corrected in advance, improving the "
             "accuracy of every invoice and reducing the manual rework the billing cycle currently generates.")

    add_heading(doc, "Selecting which tenants to review", 3, "1.4.2")
    add_body(doc,
             "A pre-billing run starts on a single selection screen where the user narrows the report by "
             "Property Portfolio, Property, Accommodation Type and Tenant, in any combination - selecting more "
             "than one of these narrows the run further rather than replacing the others, so a single property "
             "within a portfolio, or a single tenant's units within a property, can all be reviewed together or "
             "separately. At least one of Property, Property Portfolio or Tenant must be chosen; an entirely "
             "unfiltered run across every tenant in every property is not permitted, keeping each run to a "
             "manageable, deliberately scoped selection.")

    add_heading(doc, "Choosing the billing period and how far back to look", 3, "1.4.3")
    add_body(doc,
             "Alongside the tenant selection, the user picks a Latest Billing Period - defaulting to the month "
             "currently being billed - and a Number Of Periods, defaulting to four months to match the MRI "
             "pre-billing check Quorum is used to, up to a maximum of twenty-four. The report then builds one "
             "period block per month across that window for every matching tenant, each starting from its own "
             "opening balance and finishing at a closing balance carried into the next period, so a reviewer can "
             "see exactly how a tenant's account moved month by month leading up to the billing run.")

    add_heading(doc, "Itemised charges, receipts and credits with a running balance", 3, "1.4.4")
    add_body(doc,
             "Within each billing period, every invoice, payment, credit and other account movement posted "
             "against the tenant is listed with its date, document number, allocation and remarks, alongside the "
             "exclusive, tax and inclusive amounts - matching the layout finance already knows from MRI's own "
             "pre-billing check. Where an invoice carries item lines, the report replaces the single invoice row "
             "with its individual line items instead, so a reviewer can see exactly what was charged rather than "
             "just the invoice total. Each period closes with its own subtotal, carried forward as the opening "
             "balance of the next period, giving a continuous running balance across the whole window reviewed.")

    add_heading(doc, "Reconciliation and bad debt visibility", 3, "1.4.5")
    add_body(doc,
             "Two checks protect the accuracy of the report. First, if an invoice's itemised lines don't add up "
             "to the amount actually posted against the tenant's account, the difference is shown as its own "
             "\"Posted to AR but not itemised\" line rather than being silently absorbed, flagging exactly which "
             "invoices need investigating before billing. Second, every charge and settlement is classified into "
             "a movement bucket - ordinary charges, receipts, credits, bad debt written off, or bad debt "
             "recovered - so a reviewer can immediately see where a tenant's balance has been affected by a "
             "write-off or a recovery, tagged clearly against the relevant line, rather than needing to interpret "
             "it from the raw transaction alone.")

    add_heading(doc, "Portfolio-wide summary and proof", 3, "1.4.6")
    add_body(doc,
             "Once every selected tenant has been listed, the report closes with a portfolio-wide summary: total "
             "opening balance, total movement broken down the same way (charges, receipts, credits, write-offs, "
             "recoveries) grouped by allocation, and the resulting closing balance. That closing figure is proved "
             "against the opening balance plus total movement, and if the two don't agree - beyond a very small "
             "rounding tolerance - the report shows an explicit warning calling out the variance, so a reconciling "
             "issue affecting the whole selection is caught even if no single tenant's own figures looked wrong.")

    add_heading(doc, "Reviewing, printing and exporting the report", 3, "1.4.7")
    add_body(doc,
             "The finished report can be reviewed directly on screen, downloaded as a branded PDF carrying the "
             "relevant subsidiary's logo in its header, or exported as a CSV for further analysis in a "
             "spreadsheet. A very large selection - more tenants or periods than can reasonably be laid out on a "
             "PDF - is capped and the user is directed to narrow the selection or use the CSV export instead, so "
             "the report never fails silently on an oversized run; the on-screen and CSV views carry no such "
             "limit.")

    doc.add_page_break()

    # ---- 2. Components ----
    add_heading(doc, "Components", 1, "2.")
    add_body(doc,
             "This customisation introduces no new custom records, fields or lists of its own. It reads Quorum's "
             "standard AR transactions (invoices, payments, credits, deposits and journals) via SuiteQL, together "
             "with the property/lease hierarchy - Building, Unit, Utilised Charges and their supporting lists - "
             "that belongs to the separate Quorum Contracts customisation. Per that document's own assumption, "
             "its monthly scripts must already have run for the periods being reviewed, since the invoice and AR "
             "data this report reads depends on them.")

    add_heading(doc, "Custom Records Referenced (pre-existing)", 2, "2.1")
    add_table(doc, ["Record Name", "Record Internal ID", "Field(s) Used", "Purpose Here"], [
        ["Building (Property segment)", "customrecord_cseg_bb1_building", "custrecord_bb1_building_portfolio",
         "Filters the tenant directory by Property Portfolio, and identifies each tenant's property"],
        ["Unit (segment)", "customrecord_cseg_bb1_unit", "custrecord_bb1_unit_accommodation_type",
         "Filters the tenant directory by Accommodation Type, and identifies each tenant's unit"],
        ["Utilised Charges", "customrecord_bb1_utilised_charges",
         "custrecord_bb1_utilised_invoice, _item, _type, _status, _discount_amount",
         "Classifies each invoice/item combination's charge type and status, and supplies the discount amount "
         "shown against a discounted line"],
    ], [4.0, 5.0, 4.5, 3.5], mono_cols={1, 2},
       note="All three records belong to the Quorum Contracts customisation (Section 1.2) and are not created or "
            "modified by this project's SDF source.")

    add_heading(doc, "Custom Fields Referenced (pre-existing)", 2, "2.2")
    add_table(doc, ["Field Name", "Field Internal ID", "Applied To", "Description"], [
        ["Building (Property)", "cseg_bb1_building", "Transaction - Custom Segment",
         "Identifies which property a transaction belongs to; joined against for the Property/Portfolio filters "
         "and printed as the tenant's Property"],
        ["Unit", "cseg_bb1_unit", "Transaction - Custom Segment",
         "Identifies which unit a transaction belongs to; joined against for the Accommodation Type filter and "
         "printed as the tenant's Unit No."],
        ["Subsidiary Logo (Forms)", "logo", "Subsidiary (standard)",
         "Image shown in the PDF's header - loaded via record.load rather than a saved search, since this field's "
         "value isn't returned reliably by search.lookupFields"],
    ], [4.5, 4.0, 4.0, 4.5], mono_cols={1})

    add_heading(doc, "Custom Lists Referenced (pre-existing)", 2, "2.3")
    add_table(doc, ["List Name", "List Internal ID", "Purpose Here"], [
        ["Property Portfolio", "customlist_bb1_building_prop_portfolio",
         "Source list for the Property Portfolio multi-select filter"],
        ["Accommodation Type", "customlist_bb1_building_accommoda_type",
         "Source list for the Accommodation Type multi-select filter"],
    ], [5.5, 6.0, 5.5], mono_cols={1},
       note="Both lists belong to the Quorum Contracts customisation (Section 1.2) and are not created by this "
            "project's SDF source.")

    add_heading(doc, "Scripts", 2, "2.4")
    add_body(doc, "File cabinet path: SuiteScripts/[QPG] Pre-Billing/", size=9, color=BB1_GREY, space_after=10)
    add_table(doc, ["File Name", "Type", "Function", "Deployed To"], [
        ["bb1_qpg_prebill_sl.js", "su",
         "Suitelet - renders the selection criteria form (Portfolio/Property/Accommodation Type/Tenant, "
         "period range, options) and dispatches to screen, PDF or CSV rendering based on the submitted mode",
         "Standalone - script/deployment id to be confirmed at deployment (not yet tracked in this project's "
         "SDF Objects folder - see Section 3.2)"],
    ], [4.7, 1.0, 7.9, 3.4], mono_cols={0})

    add_heading(doc, "Other Files", 2, "2.5")
    add_body(doc, "Supporting library module used by the script above - not a separate deployed script in its "
                  "own right.", size=9, color=BB1_GREY, space_after=10)
    add_table(doc, ["File Name", "Function"], [
        ["bb1_qpg_prebill_lib.js", "Shared library - reads and validates the selection criteria; assembles the "
                                    "tenant directory, opening balances, AR activity, invoice lines and Utilised "
                                    "Charges classification via SuiteQL; classifies bad debt write-off/recovery; "
                                    "builds the per-tenant period ledger and portfolio-wide summary/proof; and "
                                    "renders the on-screen HTML, PDF and CSV output. Loaded by "
                                    "bb1_qpg_prebill_sl.js as a CommonJS module (./bb1_qpg_prebill_lib), not "
                                    "deployed independently"],
    ], [5.5, 11.5], mono_cols={0})

    add_heading(doc, "Entity Relationship Diagram", 2, "2.6")
    add_body(doc, "Represented as a relationship table below (no image-rendering environment was available when "
                   "this document was generated).", size=8, italic=True, color=BB1_GREY, space_after=10)
    add_table(doc, ["From", "Link Field", "To"], [
        ["Transaction (Invoice/Payment/Credit/Deposit/Journal)", "entity", "Customer (Tenant)"],
        ["Transaction", "cseg_bb1_building", "Building (Property segment)"],
        ["Transaction", "cseg_bb1_unit", "Unit (segment)"],
        ["TransactionLine / TransactionAccountingLine", "transaction / account (AcctRec)", "Transaction / Account"],
        ["TransactionLine (invoice)", "item", "Item (itemised invoice line, Allocation column)"],
        ["customrecord_bb1_utilised_charges", "custrecord_bb1_utilised_invoice / _item",
         "Transaction (Invoice) / Item - charge type/status classification and discount amount"],
        ["customrecord_cseg_bb1_building", "custrecord_bb1_building_portfolio",
         "customlist_bb1_building_prop_portfolio (Property Portfolio filter)"],
        ["customrecord_cseg_bb1_unit", "custrecord_bb1_unit_accommodation_type",
         "customlist_bb1_building_accommoda_type (Accommodation Type filter)"],
        ["TransactionAccountingLine", "account (non-AR)", "Account - offset account classification for bad debt "
                                                            "write-off/recovery, when configured"],
        ["Customer", "subsidiary", "Subsidiary - resolves the PDF header logo (Section 2.2)"],
    ], [6.0, 5.5, 5.5], mono_cols={1})

    add_heading(doc, "Additional Components", 2, "2.7")
    add_table(doc, ["Item", "Detail"], [
        ["Script parameters (classification & governance)",
         "custscript_bb1_qpg_prebill_bdwoitems_su / _bdrecitems_su (comma-separated item ids classifying a line "
         "as bad debt written off / recovered), _bdaccts_su / _bdrecaccts_su (comma-separated account ids, "
         "default 1182 / 1314, classifying a settlement by its offsetting account), _allocpfx_su (comma-"
         "separated prefixes stripped from an allocation name, default XXX_), and _maxtenant_su (caps a single "
         "run at 250 tenants by default, to stay inside NetSuite's governance limits)"],
        ["PDF row cap", "MAX_PDF_ROWS = 4,000. A selection producing more rows than this raises a clear error "
                         "directing the user to narrow the selection or use the CSV export, rather than letting "
                         "the PDF renderer fail with a generic error"],
        ["Subsidiary logo scaling", "The Subsidiary Logo (Forms) field's image is loaded via record.load (not a "
                                      "saved search, which doesn't reliably return this field), and its natural "
                                      "pixel dimensions are parsed directly from the file's GIF/PNG/JPEG bytes "
                                      "(N/file has no binary accessor) so it can be scaled down - never up - to "
                                      "fit a fixed header box while preserving its aspect ratio"],
        ["Reconciliation tolerance", "RECONCILE_TOLERANCE = 0.01. Both the per-invoice line/AR variance check and "
                                       "the portfolio-wide proof warning use this as their rounding tolerance"],
        ["Deviation from the original Service Request", "The Service Request (Section 1.2) proposed a Building > "
            "Block > Floor > Unit cascading multi-select selection screen and a PDF laid out per unit with a "
            "separate Statement Date/Start Date pair. The delivered Suitelet instead follows an MRI-style "
            "\"Tenant Billing History\" design: Property Portfolio/Property/Accommodation Type/Tenant filters "
            "(no Block/Floor level), and a single Latest Billing Period + Number Of Periods pair driving a "
            "rolling multi-period ledger, with bad debt classification and a portfolio-wide reconciliation "
            "summary the original spec did not describe. Confirm with Quorum and Andile whether this simplified, "
            "MRI-aligned filter set and layout meets the original requirement, or whether the Block/Floor "
            "cascading selection and the per-unit PDF layout described in the Service Request still need to be "
            "added"],
    ], [4.8, 12.2])

    add_heading(doc, "Process Flow", 2, "2.8")
    add_body(doc, "Represented as a swimlane-style table below (no image-rendering environment was available "
                   "when this document was generated).", size=8, italic=True, color=BB1_GREY, space_after=10)
    add_table(doc, ["Actor", "Step", "Produces"], [
        ["User", "Opens the Pre-Billing Report Suitelet URL with no Property/Portfolio/Tenant selected",
         "buildSelectionForm renders the criteria form: Property Portfolio, Property, Accommodation Type, "
         "Tenant, Latest Billing Period, Number Of Periods, and the Itemise Invoice Lines / Show Movement "
         "Analysis / Include Tenants With No Activity options"],
        ["User", "Selects at least one of Property, Property Portfolio or Tenant (optionally narrowed further "
                  "by Accommodation Type), sets the period range, clicks Run Report",
         "onRequest re-reads the filters (readFilters) and, now that a valid selection is present, proceeds to "
         "build the report rather than re-showing the criteria form"],
        ["Suitelet (lib)", "buildHistory - tenant directory & balances",
         "getTenantDirectory resolves each tenant's most recent property/unit via SuiteQL; resolveReportSubsidiary "
         "picks the header logo (warning if the selection spans more than one subsidiary); getOpeningBalances "
         "aggregates every AR posting before the reporting window per tenant"],
        ["Suitelet (lib)", "buildHistory - activity & classification",
         "getActivity pulls every AR-posting transaction inside the window; getInvoiceLines itemises invoice "
         "lines when \"Itemise Invoice Lines\" is on; getUtilisedCharges classifies each invoice/item's charge "
         "type, status and discount; decorateOffsetAccounts resolves offsetting accounts for settlements, when "
         "the write-off/recovery account parameters are configured"],
        ["Suitelet (lib)", "assembleTenants",
         "Folds any activity before the reporting window into the opening balance, builds one period block per "
         "tenant per month with a running balance, classifies each row into a movement bucket (charges/"
         "receipts/credits/write-off/recovery/other), and raises an \"Posted to AR but not itemised\" row where "
         "an invoice's lines don't reconcile to its posted amount"],
        ["Suitelet (lib)", "summariseAll",
         "Rolls every tenant into a portfolio-wide movement summary grouped by allocation, and proves Balance "
         "B/f plus total movement against Balance C/f, warning if the two don't agree"],
        ["User", "Views the on-screen HTML report, or clicks Download PDF / Download CSV",
         "renderScreen shows the report inline; renderPdf renders a branded PDF with the subsidiary logo header "
         "(capped at 4,000 rows - narrower selections only); renderCsv streams a flat CSV export with no row cap"],
    ], [3.2, 6.8, 7.0])

    doc.add_page_break()

    # ---- 3. Testing & Setup ----
    add_heading(doc, "Testing & Setup", 1, "3.")

    add_heading(doc, "Testing Details", 2, "3.1")
    add_body(doc,
              "To be completed by the tester following execution of BB1 QA testing / UAT for this "
              "customisation. Record each test case, the steps taken, the expected and actual results, and its "
              "pass/fail status below.",
              italic=True, color=BB1_GREY, space_after=10)
    add_table(doc, ["Test Area", "Test Case", "Steps", "Expected Result", "Actual Result", "Status"], [
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", "", ""],
    ], [2.3, 2.7, 4.0, 3.5, 3.0, 1.5],
       note="Suggested Test Areas to cover: selection screen validation (at least one of Property/Portfolio/"
            "Tenant required); combined filters (Portfolio + Property + Accommodation Type + Tenant together); "
            "Number Of Periods boundaries (1, the default 4, and the maximum 24); the itemised-invoice-line vs. "
            "single-invoice-row toggle; the invoice line/AR reconciliation variance row; bad debt classification "
            "by item id and by offset account; the portfolio-wide proof warning when it doesn't balance; "
            "subsidiary logo rendering and its missing/inaccessible fallback; a multi-subsidiary selection's "
            "warning; the PDF row cap (over 4,000 rows) directing the user to CSV; and the CSV export's content "
            "against the same selection's on-screen figures.")

    add_heading(doc, "Deployment & Configuration Setup", 2, "3.2")
    add_body(doc,
              "This customisation creates no new custom records, fields or lists, but the following deployment-"
              "time steps are required and are not captured in this project's SDF source, since Script and "
              "Script Deployment objects for this Suitelet are not yet tracked in the Objects folder.",
              space_after=8)
    add_numbered_list(doc, [
        "Deploy bb1_qpg_prebill_sl.js as a Suitelet (customscript_bb1_qpg_prebill_su, exact id to be confirmed), "
        "with bb1_qpg_prebill_lib.js present in the same file cabinet folder so the Suitelet's require('./"
        "bb1_qpg_prebill_lib') resolves.",
        "Add Free-Form Text Script Parameters to the deployment: custscript_bb1_qpg_prebill_bdwoitems_su, "
        "custscript_bb1_qpg_prebill_bdrecitems_su, custscript_bb1_qpg_prebill_bdaccts_su (default 1182), "
        "custscript_bb1_qpg_prebill_bdrecaccts_su (default 1314), custscript_bb1_qpg_prebill_allocpfx_su "
        "(default XXX_), and custscript_bb1_qpg_prebill_maxtenant_su (default 250) - each a comma-separated "
        "list where the parameter expects one.",
        "Confirm the pre-existing customrecord_cseg_bb1_building, customrecord_cseg_bb1_unit and "
        "customrecord_bb1_utilised_charges records, and the customlist_bb1_building_prop_portfolio / "
        "customlist_bb1_building_accommoda_type lists (Quorum Contracts customisation, Section 2.1/2.3) already "
        "exist and are populated.",
        "Confirm the Subsidiary Logo (Forms) field is populated with a JPG, GIF or PNG image on every subsidiary "
        "that will appear in a pre-billing run; a missing or inaccessible logo degrades gracefully to a header "
        "without one, surfaced as a report warning rather than a failure.",
        "Confirm the Quorum Contracts monthly scripts have already run for the periods being reviewed, per the "
        "Service Request's stated assumption (Section 1.2), since the invoice and AR data this report reads "
        "depends on them.",
        "Assign role/permission access to the Suitelet for the finance/billing team who will run the Pre-Billing "
        "Report ahead of each monthly billing run.",
        "Confirm with Quorum and Andile whether the delivered filter set and MRI-style layout meet the original "
        "Service Request, or whether the Block/Floor cascading selection and per-unit PDF layout it originally "
        "described still need to be added (Section 2.7).",
    ])

    doc.save(OUTPUT_PATH)
    print(f"Written: {OUTPUT_PATH}")


if __name__ == "__main__":
    build()
