"""
BB1 Technical Notes generator - J&K Ross, SO Pick Ticket (Gap 9).

Built with python-docx (no Node/npm available on this machine, so this adapts the
standard BB1 docx-package recipe to python-docx while keeping the same brand rules:
colours, table styling, DXA-fixed column widths, and zero-width-space soft breaks on
long NetSuite identifiers so they don't bleed off the page edge.

Usage: python3 build_docx.py
Output: ../BB1_TechNotes_JKR_SOPickTicket.docx (relative to this script)
"""
import os
import re
import docx
from docx import Document
from docx.shared import Twips, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "..", "BB1_TechNotes_JKR_SOPickTicket.docx")

BB1_BLUE = RGBColor(0x24, 0x78, 0xAA)
BB1_RED = RGBColor(0xC0, 0x28, 0x2E)
BB1_GREY = RGBColor(0x9B, 0x96, 0x92)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
MONO_GREY = RGBColor(0x44, 0x44, 0x44)

FONT = "Roboto"
MONO_FONT = "Courier New"

CONTENT_WIDTH = 9638  # DXA

ZWSP = "​"


def soft_break(s):
    """Insert a zero-width space after every underscore/dot in long tokens so Word
    can wrap NetSuite identifiers instead of letting them bleed off the page."""
    if s is None:
        return ""
    s = str(s)
    tokens = re.split(r"(\s+)", s)
    out = []
    for tok in tokens:
        if len(tok) >= 15:
            tok = re.sub(r"([_.])", lambda m: m.group(1) + ZWSP, tok)
        out.append(tok)
    return "".join(out)


# ---------------------------------------------------------------------------
# Style setup
# ---------------------------------------------------------------------------

def setup_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BB1_BLUE

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle.font.size = Pt(18)
    subtitle.font.bold = True
    subtitle.font.color.rgb = BB1_BLUE
    subtitle.font.italic = False

    h1 = styles["Heading 1"]
    h1.font.name = FONT
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BB1_BLUE

    h2 = styles["Heading 2"]
    h2.font.name = FONT
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = BB1_RED

    h3 = styles["Heading 3"]
    h3.font.name = FONT
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.color.rgb = BB1_BLUE

    # section margins (2cm ~= 1134 twips), A4 page size
    for section in doc.sections:
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
        section.top_margin = Twips(1134)
        section.bottom_margin = Twips(1134)
        section.left_margin = Twips(1134)
        section.right_margin = Twips(1134)


def add_grey_note(doc, text, italic=True, size=9):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = BB1_GREY
    r.font.italic = italic
    return p


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.bold = bold
    p.paragraph_format.space_after = Pt(8)
    return p


def add_numbered(doc, items, style="List Number"):
    for item in items:
        p = doc.add_paragraph(style=style)
        r = p.add_run(item)
        r.font.name = FONT
        r.font.size = Pt(10)


# ---------------------------------------------------------------------------
# Table helpers
# ---------------------------------------------------------------------------

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def set_cell_width(cell, twips):
    cell.width = Twips(twips)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def set_table_grid(table, col_widths):
    """python-docx creates an equal-width tblGrid on add_table(); Word renders columns
    from this grid regardless of per-cell tcW unless it's overwritten explicitly."""
    tbl = table._tbl
    tblGrid = tbl.find(qn("w:tblGrid"))
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in col_widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        tblGrid.append(gc)


def set_table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CCCCCC")
        borders.append(el)
    tblPr.append(borders)


def set_cell_text(cell, value, mono=False, bold=False, color=None, size=10, wrap_ids=True):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    values = value if isinstance(value, (list, tuple)) else [value]
    for i, v in enumerate(values):
        target_p = p if i == 0 else cell.add_paragraph()
        target_p.paragraph_format.space_after = Pt(0)
        txt = soft_break(v) if wrap_ids else str(v)
        r = target_p.add_run(txt)
        r.font.name = MONO_FONT if mono else FONT
        r.font.size = Pt(9 if mono else size)
        r.font.bold = bold
        r.font.color.rgb = color if color else (MONO_GREY if mono else BLACK)


def add_styled_table(doc, headers, rows, col_widths, mono_cols=None, zebra=True):
    """headers: list[str]; rows: list[list[str|list[str]]]; col_widths: list[int] (DXA).
    mono_cols: set of column indexes to render in monospace with ID-safe soft breaks."""
    mono_cols = mono_cols or set()
    assert sum(col_widths) == CONTENT_WIDTH, f"column widths must sum to {CONTENT_WIDTH}, got {sum(col_widths)}"

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, col_widths)
    set_table_fixed_layout(table)
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_width(hdr_cells[i], col_widths[i])
        set_cell_margins(hdr_cells[i])
        set_cell_shading(hdr_cells[i], "2478AA")
        set_cell_text(hdr_cells[i], h, bold=True, color=WHITE, size=11, wrap_ids=False)

    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "F5F5F5" if (zebra and r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            set_cell_width(cells[c_idx], col_widths[c_idx])
            set_cell_margins(cells[c_idx])
            set_cell_shading(cells[c_idx], fill)
            set_cell_text(cells[c_idx], val, mono=(c_idx in mono_cols))
    return table


def add_image(doc, filename, width_in, caption=None):
    path = os.path.join(ASSETS, filename)
    from PIL import Image as PILImage
    with PILImage.open(path) as im:
        w, h = im.size
    height_in = width_in * (h / w)
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.font.name = FONT
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = BB1_GREY
    return height_in


def add_page_break(doc):
    doc.add_page_break()


# ---------------------------------------------------------------------------
# Document build
# ---------------------------------------------------------------------------

def build():
    doc = Document()
    setup_styles(doc)

    # ---------------- Cover ----------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("BB1")
    r.font.name = FONT
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    sub = doc.add_paragraph()
    r = sub.add_run("Technical Notes")
    r.font.name = FONT
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    cust = doc.add_paragraph()
    r = cust.add_run("J&K Ross")
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = BB1_RED

    title = doc.add_paragraph()
    r = title.add_run("Automated Printing - Sales Order Pick Ticket (Gap 9)")
    r.font.name = FONT
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    meta = doc.add_paragraph()
    r = meta.add_run("Author: Jared Espineli   |   Date: 18 September 2026")
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.color.rgb = BB1_GREY

    add_grey_note(
        doc,
        "Status: Solution design - pending sign-off. See Section 4, Action Items, "
        "before any build work begins.",
    )

    add_page_break(doc)

    # ---------------- 1. Overview ----------------
    doc.add_heading("1. Overview", level=1)

    doc.add_heading("1.1 Introduction", level=2)
    add_body(
        doc,
        "J&K Ross needs NetSuite to print operational warehouse documents automatically, "
        "without a user manually opening and printing a PDF. This document covers the first "
        "and most fully designed use case in Gap 9 of the functional design: printing the "
        "Sales Order Pick Ticket the moment an order is released for fulfilment and inventory "
        "has been committed against it. The same underlying engine - a reusable SuiteFlow "
        "Workflow Action Script paired with a client-side QZ Tray bridge - is intended to be "
        "reused for the other document types in the functional design (Work Order Pick Note, "
        "Project Order Picking Ticket, Manufacturing Document, Address Labels), but those rows "
        "are only sketched at the matrix level today and are out of scope for this revision.",
    )

    doc.add_heading("1.2 BRS / Specification", level=2)
    add_body(
        doc,
        "Functional Design: Gap 9 (BlueBridge One), source file "
        "“P102821-Functional Design_ Gap 9-180926-102113.pdf”, provided directly by the "
        "user for this project. No Google Drive link was supplied for this revision.",
    )

    doc.add_heading("1.3 Development Environment", level=2)
    add_body(doc, "Production Account: 11643893")

    doc.add_heading("1.4 Business Requirements", level=2)

    doc.add_heading("1.4.1 Summary", level=3)
    add_body(
        doc,
        "Warehouse staff currently have to remember to print a pick ticket every time an "
        "order is ready to be picked, which is repetitive and occasionally missed. This "
        "customisation removes that manual step for Sales Orders: as soon as an order moves "
        "into an open fulfilment state and stock has been committed to it, NetSuite "
        "automatically renders the pick ticket and sends it straight to the correct printer "
        "in the warehouse, with no one needing to click Print. The result is that pick tickets "
        "are produced consistently and at the right moment, freeing staff to focus on picking "
        "and packing rather than paperwork, and reducing the risk of an order being missed "
        "because its ticket was never printed.",
    )

    doc.add_heading("1.4.2 Automated Pick Ticket Printing", level=3)
    add_body(
        doc,
        "When a Sales Order is saved and its status reaches an open fulfilment state such as "
        "Pending Fulfillment, with at least one line showing committed quantity, the system "
        "recognises that the order is ready to be picked. At that point a print request for "
        "the Sales Order Pick Ticket is raised automatically and one copy is queued for "
        "delivery to the warehouse printer assigned to that order's location - no user action "
        "is required to trigger it. Determining exactly when a print should fire, and for "
        "which orders, remains the responsibility of the NetSuite workflow rather than the "
        "print engine itself, so the trigger conditions can be tuned for J&K Ross without "
        "touching the reusable script.",
    )

    doc.add_heading("1.4.3 Printer Configuration and Resolution", level=3)
    add_body(
        doc,
        "Each warehouse printer is registered once in NetSuite with its name and network "
        "address, matching how it is set up in the QZ Tray printing software on the warehouse "
        "computer. Every warehouse Location is then linked to the printer that should handle "
        "its pick tickets. This means the system always knows which physical printer to send a "
        "document to, based purely on where the order is being fulfilled from, and a new "
        "printer can be added or reassigned by an administrator without any script changes.",
    )

    doc.add_heading("1.4.4 Print Job Queue and Client-Side Delivery", level=3)
    add_body(
        doc,
        "Because the decision to print happens on NetSuite's server the moment an order is "
        "saved, but the actual printing has to happen from the warehouse computer's browser "
        "(where the QZ Tray printing software is running), the system holds each print "
        "request briefly in a queue. As soon as the user's browser returns to the order after "
        "saving, it checks the queue for anything waiting against that order, retrieves the "
        "ready-to-print document, and delivers it to the printer. Once delivery is confirmed, "
        "the queued request is marked as complete so the same ticket is never printed twice, "
        "for example if the page is refreshed.",
    )

    doc.add_heading("1.4.5 Reusability and Future Extension", level=3)
    add_body(
        doc,
        "The print engine itself does not know anything specific about Sales Orders, Pick "
        "Tickets, or J&K Ross - it simply receives a document type, a template, a printer, "
        "and a quantity, and takes care of rendering and delivery. This means the same "
        "building blocks used here can be reused to automatically print Work Order pick notes, "
        "manufacturing documents, or address labels, and can be configured for other "
        "customers entirely, simply by adding new workflow configuration rather than writing "
        "new print logic.",
    )

    add_page_break(doc)

    # ---------------- 2. Components ----------------
    doc.add_heading("2. Components", level=1)

    # 2.1 Custom Records
    doc.add_heading("2.1 Custom Records", level=2)
    add_styled_table(
        doc,
        headers=["#", "Record Name", "Record Internal ID", "Function"],
        rows=[
            ["1", "Printer",
             "customrecord_bb1_printer",
             "Stores each warehouse printer's name and IP address so the print engine can "
             "hand off the right payload to QZ Tray."],
            ["2", "Print Job Queue",
             "customrecord_bb1_print_job",
             "Bridges the server-side After Submit trigger and the client-side QZ Tray "
             "delivery. Holds a print request as Pending until a browser with QZ Tray "
             "running can render and deliver it, then tracks it through Sent, Printed or "
             "Failed."],
        ],
        col_widths=[500, 2500, 3100, 3538],
        mono_cols={2},
    )

    doc.add_heading("customrecord_bb1_printer - Fields", level=3)
    add_styled_table(
        doc,
        headers=["Field Name", "Field Internal ID", "Type", "Description"],
        rows=[
            ["Printer Name", "_bb1_prn_name", "Free Text",
             "Descriptive name of the warehouse printer, matching how it is registered in "
             "QZ Tray on the warehouse machine."],
            ["IP Address", "_bb1_prn_ip_address", "Free Text",
             "Network IP address of the printer, used when sending raw ZPL commands over a "
             "socket instead of a rendered PDF."],
            ["Printer Type (recommended new field)", "_bb1_prn_type", "List",
             "PDF/Raster vs Raw/ZPL. Not in the original FDD; needed so the Suitelet knows "
             "how to build the QZ Tray payload for this printer."],
        ],
        col_widths=[2500, 3000, 1600, 2538],
        mono_cols={1},
    )

    doc.add_heading("customrecord_bb1_print_job - Fields", level=3)
    add_styled_table(
        doc,
        headers=["Field Name", "Field Internal ID", "Type", "Description"],
        rows=[
            ["Source Record", "custrecord_bb1_pj_source_rec", "Record Reference",
             "Internal ID of the transaction that triggered the print, e.g. the Sales Order."],
            ["Source Record Type", "custrecord_bb1_pj_source_type", "Free Text",
             "Record type of the source, e.g. salesorder, kept generic so the same queue can "
             "later serve Work Order, Project Order and Item Fulfillment."],
            ["Document Type", "custrecord_bb1_pj_doctype", "Free Text",
             "e.g. “Pick Ticket”, passed straight through from the workflow action "
             "parameter."],
            ["Template", "custrecord_bb1_pj_template", "Free Text",
             "Internal ID of the Advanced PDF/HTML template used to render the document "
             "(owned by Gap 2 / Gap 7)."],
            ["Printer", "custrecord_bb1_pj_printer", "List/Record",
             "Target printer resolved at trigger time; links to customrecord_bb1_printer."],
            ["Quantity", "custrecord_bb1_pj_qty", "Integer",
             "Number of copies to print - 1 for a Pick Ticket, driven by Number of Boxes for "
             "Address Labels."],
            ["Status", "custrecord_bb1_pj_status", "List",
             "Pending / Sent / Printed / Failed - stops a page refresh from triggering a "
             "duplicate print."],
        ],
        col_widths=[2500, 3000, 1600, 2538],
        mono_cols={1},
    )

    # 2.2 Custom Fields
    doc.add_heading("2.2 Custom Fields", level=2)
    add_styled_table(
        doc,
        headers=["Field Name", "Field Internal ID", "Type", "Applied To", "Source / Default"],
        rows=[
            ["Default Printer", "custrecord_bb1_loc_default_printer", "List/Record",
             "Location",
             "Set manually per warehouse Location; read by the workflow action to resolve "
             "which printer a Sales Order's Pick Ticket should print to."],
        ],
        col_widths=[2000, 2600, 1550, 1550, 1938],
        mono_cols={1},
    )

    # 2.4 Scripts
    doc.add_heading("2.4 Scripts", level=2)
    add_grey_note(
        doc,
        "File cabinet path: src/FileCabinet/SuiteScripts/[JKR] Print Document/ - "
        "bb1_jkr_prnt_doc_wfa.js currently exists in the repository as an empty stub "
        "(WorkflowActionScript shell, onAction not yet implemented). "
        "bb1_print_bridge_sl.js and bb1_print_bridge_cs.js do not exist yet.",
    )
    add_styled_table(
        doc,
        headers=["File Name", "Type", "Function", "Deployed To"],
        rows=[
            ["bb1_jkr_prnt_doc_wfa.js", "wfa",
             "Reads workflow action parameters (Document Type, Printer, Template, Quantity) "
             "from the triggering workflow and creates a customrecord_bb1_print_job record "
             "with status Pending. Deliberately does not render the PDF or talk to QZ Tray, "
             "keeping the engine generic and reusable across document types.",
             "Sales Order - After Submit action on the “BB1 - Auto Print” workflow"],
            ["bb1_print_bridge_sl.js", "su",
             "GET endpoint that finds Pending jobs for a given source record, renders the "
             "PDF via N/render using the job's template, base64-encodes it, returns the "
             "payload plus target printer name/IP/type as JSON, and marks the job Sent. Also "
             "accepts an acknowledgement call to mark a job Printed or Failed.",
             "Standalone Suitelet, called by bb1_print_bridge_cs.js"],
            ["bb1_print_bridge_cs.js", "cs",
             "pageInit client script that calls the Suitelet with the current record's id "
             "and type. Where a pending job is returned, it uses the qz-tray.js library "
             "(existing QZ integration by Wayne Wundram) to open a WebSocket connection to "
             "the local QZ Tray instance and deliver the payload, then acknowledges the job.",
             "Sales Order (and, as further document types are added, Work Order / Item "
             "Fulfillment / Project Order)"],
        ],
        col_widths=[2400, 1000, 3600, 2638],
        mono_cols={0},
    )

    # 2.5 Other Files
    doc.add_heading("2.5 Other Files", level=2)
    add_styled_table(
        doc,
        headers=["File Name", "Function"],
        rows=[
            ["qz-tray.js",
             "Third-party QZ Industries client library, wrapped by Wayne Wundram's existing "
             "BB1 QZ library, that opens a WebSocket connection from the browser to the "
             "locally installed QZ Tray application and dispatches the print payload to the "
             "OS-registered printer. Not a NetSuite script - loaded by the client script from "
             "the File Cabinet or a trusted CDN."],
        ],
        col_widths=[3200, 6438],
        mono_cols={0},
    )

    # 2.6 ERD
    doc.add_heading("2.6 ERD", level=2)
    add_body(
        doc,
        "The diagram below shows how the new Printer and Print Job Queue records relate to "
        "the standard Sales Order and Location records. The Source Record link on the queue "
        "is deliberately generic (polymorphic) so the same record can later hold jobs raised "
        "from other transaction types.",
    )
    add_image(doc, "erd.png", width_in=6.3)

    # 2.7 Additional Components
    doc.add_heading("2.7 Additional Components", level=2)
    doc.add_heading("SuiteFlow Workflow - \"BB1 - Auto Print\" (Sales Order)", level=3)
    add_body(doc, "Trigger: After Submit", bold=False)
    add_numbered(doc, [
        "Condition: Order Status is Pending Fulfillment (extend to other open statuses as "
        "confirmed) AND a body-level Formula (Numeric) field evaluating "
        "MAX({quantitycommitted}) > 0 across the item sublist.",
        "Action: custom workflow action “BB1 Print Document”, calling "
        "bb1_jkr_prnt_doc_wfa.js, with parameters: Document Type = “Pick Ticket” "
        "(constant); Quantity = 1 (constant); Printer = {location.custrecord_bb1_loc_default_"
        "printer} (field lookup); Template = internal ID of the SO Pick Ticket Advanced "
        "PDF/HTML template (Gap 2 / Gap 7 deliverable, not yet confirmed - see Action Items).",
    ])

    # 2.8 Process Flow
    doc.add_heading("2.8 Process Flow", level=2)
    add_body(
        doc,
        "The flow below answers the functional design's Open Point #1 (how the server-side "
        "action script hands the print payload to the browser): rather than pushing from the "
        "server, the action script only queues a job, and a client script plus Suitelet pair "
        "picks it up the moment the user's browser returns to the record.",
    )
    add_image(doc, "flow.png", width_in=6.3)

    add_page_break(doc)

    # ---------------- 3. Testing & Setup ----------------
    doc.add_heading("3. Testing & Setup", level=1)

    doc.add_heading("3.1 Steps to Test / Replicate", level=2)
    add_numbered(doc, [
        "Create a Printer record for a test warehouse printer (Name, IP Address, Printer "
        "Type) and confirm it is registered in QZ Tray on the test machine.",
        "Set that Printer record as the Default Printer on the test Location.",
        "Ensure QZ Tray is running on the test machine and the browser has approved the QZ "
        "certificate/connection.",
        "Create a Sales Order at that Location, add a line, and commit at least one unit of "
        "inventory (Quantity Committed > 0).",
        "Save the order so its status reaches Pending Fulfillment and the “BB1 - Auto "
        "Print” workflow's After Submit condition is met.",
        "Confirm a customrecord_bb1_print_job record is created with status Pending, and "
        "the correct source record, printer, template and quantity.",
        "Confirm the browser, after redirecting to the Sales Order view, automatically "
        "prints the Pick Ticket to the assigned printer with no manual interaction, and the "
        "job status updates to Printed.",
        "Refresh the Sales Order page and confirm the ticket does not print a second time.",
        "Repeat with Quantity Committed = 0 on all lines and confirm no job is created "
        "(negative test).",
        "Repeat with no Default Printer set on the Location and confirm the job fails or "
        "logs an error rather than silently doing nothing.",
        "Simulate a headless save (e.g. via CSV import) and confirm the resulting Pending "
        "job is visible for manual follow-up, since no browser is present to complete the "
        "print.",
    ])

    doc.add_heading("3.2 Custom Record Setup", level=2)
    add_numbered(doc, [
        "Deploy customrecord_bb1_printer with fields _bb1_prn_name, _bb1_prn_ip_address, and "
        "the recommended _bb1_prn_type.",
        "Deploy customrecord_bb1_print_job with the fields listed in Section 2.1, and a list "
        "view or saved search filtered to Status = Pending for support staff to monitor.",
        "Add custrecord_bb1_loc_default_printer to the Location record and populate it for "
        "every warehouse Location that will use automated printing.",
        "Create one customrecord_bb1_printer record per physical warehouse printer, matching "
        "the printer name exactly as registered in QZ Tray on the relevant machine.",
        "Build the “BB1 - Auto Print” workflow on the Sales Order record per Section "
        "2.7.",
        "Deploy the three scripts in Section 2.4 and confirm the Client Script deployment "
        "covers every role/user that actions Sales Order fulfilment in the warehouse.",
    ])

    add_page_break(doc)

    # ---------------- 4. Action Items ----------------
    doc.add_heading("4. Action Items", level=1)
    add_body(
        doc,
        "Items 1-6 are decisions or dependencies that need to be resolved before build starts "
        "on the affected component. Items 7-17 are the build-out tasks themselves, in the "
        "order they are needed.",
    )
    add_styled_table(
        doc,
        headers=["#", "Action Item", "Category", "Notes"],
        rows=[
            ["1", "Confirm printer-to-order resolution approach", "Decision Needed",
             "Design assumes one Default Printer per Location; confirm this is right vs. "
             "per-user or per-role, especially where a warehouse has more than one pick "
             "ticket printer."],
            ["2", "Obtain the SO Pick Ticket Advanced PDF/HTML template ID", "Dependency",
             "Owned by Gap 2 / Gap 7. The workflow action's Template parameter cannot be "
             "finalised without it."],
            ["3", "Sign off the queue-and-poll bridge mechanism", "Decision Needed",
             "Answers the FDD's Open Point #1 (server-to-client handoff). Review with the "
             "technical lead and, for the qz-tray.js API/version specifics, Wayne Wundram "
             "(owner of the existing BB1 QZ library)."],
            ["4", "Approve the net-new customrecord_bb1_print_job record", "Decision Needed",
             "Not documented in the original FDD's Custom Records section - it is an "
             "addition this design introduces to solve Open Point #1, so it needs sign-off "
             "before build."],
            ["5", "Decide on a fallback for headless triggers", "Decision Needed",
             "CSV import, RESTlet or Suitelet mass-update saves have no browser to poll for "
             "the queued job, so the print never fires. Decide whether a “Pending "
             "Prints” dashboard/reminder is needed, or this is an accepted limitation."],
            ["6", "Confirm duplicate-print prevention is sufficient", "Decision Needed",
             "The Status field on the job record stops a page refresh from reprinting; "
             "confirm this is enough, and whether a manual “reprint” action is "
             "also wanted."],
            ["7", "Create customrecord_bb1_printer (+ Printer Type field)", "Build",
             "See Section 3.2, step 1."],
            ["8", "Create customrecord_bb1_print_job", "Build",
             "See Section 3.2, step 2."],
            ["9", "Add Default Printer field to Location", "Build",
             "See Section 3.2, step 3."],
            ["10", "Populate Printer records and Location defaults for all warehouse "
             "printers", "Build / Deploy",
             "See Section 3.2, steps 3-4; depends on IT confirming QZ Tray printer names/IPs "
             "per machine."],
            ["11", "Implement bb1_jkr_prnt_doc_wfa.js", "Build",
             "Currently an empty stub in the repository; implement per Section 2.4."],
            ["12", "Implement bb1_print_bridge_sl.js", "Build",
             "New script; implement per Section 2.4."],
            ["13", "Implement bb1_print_bridge_cs.js + qz-tray.js integration", "Build",
             "New script; confirm the exact qz-tray.js version/API with Wayne Wundram before "
             "coding against it (links to item 3)."],
            ["14", "Build the “BB1 - Auto Print” SuiteFlow workflow", "Build",
             "Per Section 2.7, including the Formula (Numeric) condition field."],
            ["15", "Execute the test plan", "Test",
             "Per Section 3.1, including the negative and headless-trigger cases."],
            ["16", "Confirm QZ Tray + qz-tray.js are installed and running on warehouse PCs, "
             "with printers registered under matching names", "Deploy",
             "Prerequisite for any testing beyond unit level; IT-owned."],
            ["17", "Migrate to Production (account 11643893) after sign-off", "Deploy",
             "Standard BB1 release process; do not migrate until items 1-6 are resolved."],
        ],
        col_widths=[500, 2400, 1600, 5138],
        mono_cols=set(),
    )

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
