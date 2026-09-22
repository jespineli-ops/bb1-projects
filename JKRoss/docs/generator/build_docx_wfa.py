"""
BB1 Technical Notes generator - J&K Ross, Automatic Print Documents WorkflowActionScript
(Gap 9).

Scope: this document covers only the WorkflowActionScript deliverable actually built
(bb1_jkr_prnt_doc_wfa.js). The SuiteFlow workflows that call it, and any custom records
referenced by the Functional Design, are being built by another developer and are
documented here as dependencies, not as part of this component.

Built with python-docx (no Node/npm available on this machine), following the same BB1
brand rules used in build_docx.py: colours, table styling, DXA-fixed column widths, and
zero-width-space soft breaks on long NetSuite identifiers so they don't bleed off the
page edge.

Usage: python3 build_docx_wfa.py
Output: ../BB1_TechNotes_JKR_AutoPrintWFA.docx (relative to this script)
"""
import os
import re
import docx
from docx import Document
from docx.shared import Twips, Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
OUT = os.path.join(HERE, "..", "BB1_TechNotes_JKR_AutoPrintWFA.docx")

BB1_BLUE = RGBColor(0x24, 0x78, 0xAA)
BB1_RED = RGBColor(0xC0, 0x28, 0x2E)
BB1_GREY = RGBColor(0x9B, 0x96, 0x92)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
MONO_GREY = RGBColor(0x44, 0x44, 0x44)

FONT = "Roboto"
MONO_FONT = "Courier New"

CONTENT_WIDTH = 9638  # DXA

ZWSP = "​"


def soft_break(s):
    if s is None:
        return ""
    s = str(s)
    tokens = re.split(r"(\s+)", s)
    out = []
    for tok in tokens:
        if len(tok) >= 15:
            tok = re.sub(r"([_.])", lambda m: m.group(1) + ZWSP, tok)
        out.append(tok)
    return "".join(out)


# ---------------------------------------------------------------------------
# Style setup
# ---------------------------------------------------------------------------

def setup_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(6)

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = BB1_BLUE

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle.font.size = Pt(18)
    subtitle.font.bold = True
    subtitle.font.color.rgb = BB1_BLUE
    subtitle.font.italic = False

    h1 = styles["Heading 1"]
    h1.font.name = FONT
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BB1_BLUE

    h2 = styles["Heading 2"]
    h2.font.name = FONT
    h2.font.size = Pt(14)
    h2.font.bold = True
    h2.font.color.rgb = BB1_RED

    h3 = styles["Heading 3"]
    h3.font.name = FONT
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.color.rgb = BB1_BLUE

    for section in doc.sections:
        section.page_width = Twips(11906)
        section.page_height = Twips(16838)
        section.top_margin = Twips(1134)
        section.bottom_margin = Twips(1134)
        section.left_margin = Twips(1134)
        section.right_margin = Twips(1134)


def add_grey_note(doc, text, italic=True, size=9):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.color.rgb = BB1_GREY
    r.font.italic = italic
    return p


def add_body(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.bold = bold
    p.paragraph_format.space_after = Pt(8)
    return p


def add_numbered(doc, items, style="List Number"):
    for item in items:
        p = doc.add_paragraph(style=style)
        r = p.add_run(item)
        r.font.name = FONT
        r.font.size = Pt(10)


# ---------------------------------------------------------------------------
# Table helpers
# ---------------------------------------------------------------------------

def set_cell_shading(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def set_cell_width(cell, twips):
    cell.width = Twips(twips)
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def set_table_grid(table, col_widths):
    tbl = table._tbl
    tblGrid = tbl.find(qn("w:tblGrid"))
    for gc in list(tblGrid):
        tblGrid.remove(gc)
    for w in col_widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        tblGrid.append(gc)


def set_table_fixed_layout(table):
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "CCCCCC")
        borders.append(el)
    tblPr.append(borders)


def set_cell_text(cell, value, mono=False, bold=False, color=None, size=10, wrap_ids=True):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    values = value if isinstance(value, (list, tuple)) else [value]
    for i, v in enumerate(values):
        target_p = p if i == 0 else cell.add_paragraph()
        target_p.paragraph_format.space_after = Pt(0)
        txt = soft_break(v) if wrap_ids else str(v)
        r = target_p.add_run(txt)
        r.font.name = MONO_FONT if mono else FONT
        r.font.size = Pt(9 if mono else size)
        r.font.bold = bold
        r.font.color.rgb = color if color else (MONO_GREY if mono else BLACK)


def add_styled_table(doc, headers, rows, col_widths, mono_cols=None, zebra=True):
    mono_cols = mono_cols or set()
    assert sum(col_widths) == CONTENT_WIDTH, f"column widths must sum to {CONTENT_WIDTH}, got {sum(col_widths)}"

    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid(table, col_widths)
    set_table_fixed_layout(table)
    set_table_borders(table)

    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_width(hdr_cells[i], col_widths[i])
        set_cell_margins(hdr_cells[i])
        set_cell_shading(hdr_cells[i], "2478AA")
        set_cell_text(hdr_cells[i], h, bold=True, color=WHITE, size=11, wrap_ids=False)

    for r_idx, row in enumerate(rows):
        cells = table.add_row().cells
        fill = "F5F5F5" if (zebra and r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row):
            set_cell_width(cells[c_idx], col_widths[c_idx])
            set_cell_margins(cells[c_idx])
            set_cell_shading(cells[c_idx], fill)
            set_cell_text(cells[c_idx], val, mono=(c_idx in mono_cols))
    return table


def add_image(doc, filename, width_in, caption=None):
    path = os.path.join(ASSETS, filename)
    from PIL import Image as PILImage
    with PILImage.open(path) as im:
        w, h = im.size
    height_in = width_in * (h / w)
    doc.add_picture(path, width=Inches(width_in))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.font.name = FONT
        r.font.size = Pt(9)
        r.font.italic = True
        r.font.color.rgb = BB1_GREY
    return height_in


def add_page_break(doc):
    doc.add_page_break()


# ---------------------------------------------------------------------------
# Document build
# ---------------------------------------------------------------------------

def build():
    doc = Document()
    setup_styles(doc)

    # ---------------- Cover ----------------
    p = doc.add_paragraph()
    r = p.add_run("BB1")
    r.font.name = FONT
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    sub = doc.add_paragraph()
    r = sub.add_run("Technical Notes")
    r.font.name = FONT
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    cust = doc.add_paragraph()
    r = cust.add_run("J&K Ross")
    r.font.name = FONT
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = BB1_RED

    title = doc.add_paragraph()
    r = title.add_run("Automatic Print Documents - Workflow Action Script (Gap 9)")
    r.font.name = FONT
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = BB1_BLUE

    meta = doc.add_paragraph()
    r = meta.add_run("Author: Jared Espineli   |   Date: 22 September 2026")
    r.font.name = FONT
    r.font.size = Pt(10)
    r.font.color.rgb = BB1_GREY

    add_grey_note(
        doc,
        "Status: Initial technical notes, first pass. Scope is the WorkflowActionScript "
        "component only (bb1_jkr_prnt_doc_wfa.js). The SuiteFlow workflows that call this "
        "action, and any custom records they depend on, are being built by another "
        "developer and are out of scope of this document - see Section 2.7 and Section 4.",
    )

    add_page_break(doc)

    # ---------------- 1. Overview ----------------
    doc.add_heading("1. Overview", level=1)

    doc.add_heading("1.1 Introduction", level=2)
    add_body(
        doc,
        "J&K Ross needs NetSuite to print warehouse documents - pick tickets, pick notes, "
        "labels - automatically, without a user manually opening and printing a PDF. Gap 9 "
        "of the functional design proposes a single, reusable NetSuite SuiteFlow Workflow "
        "Action Script that acts as a configurable print engine: NetSuite workflows decide "
        "when a document should print, and the action script renders and delivers it to a "
        "warehouse printer via the existing BB1 QZ Tray library. This document covers only "
        "the piece of that design actually assigned to and built by this developer - the "
        "action script itself, bb1_jkr_prnt_doc_wfa.js. The SuiteFlow workflows that will "
        "call it (one per triggering scenario in the functional design's conditions matrix, "
        "for example a Work Order/Build being saved with its finished goods committed to an "
        "open Sales Order) are being configured separately by another developer, and any "
        "custom records those workflows may need are theirs to build, not this script's.",
    )

    doc.add_heading("1.2 BRS / Specification", level=2)
    add_body(
        doc,
        "Two PDF documents were provided directly by the user for this project, with no "
        "Confluence or Drive link supplied:",
    )
    add_numbered(doc, [
        "Gap 9: Automatic Print Documents - business requirement and gap definition, "
        "including the August 2026 expansion notes on picking-slip and inventory "
        "commitment scenarios (file: P102821-Gap 9_ Automatic Print Documents-170926-"
        "100913.pdf).",
        "Functional Design: Gap 9 - architecture, action script parameters, the workflow "
        "triggers and conditions matrix, and the proposed Printer custom record (file: "
        "P102821-Functional Design_ Gap 9-180926-102113.pdf).",
    ])

    doc.add_heading("1.3 Development Environment", level=2)
    add_body(doc, "Production Account: 11643893")

    doc.add_heading("1.4 Business Requirements", level=2)

    doc.add_heading("1.4.1 Summary", level=3)
    add_body(
        doc,
        "Warehouse staff at J&K Ross currently rely on someone remembering to print a pick "
        "ticket, pick note or label at the right moment in the fulfilment or production "
        "process, which is repetitive and occasionally missed. The business wants NetSuite "
        "to do this automatically: the moment a qualifying event happens - an order being "
        "released, stock being committed, a build being completed - the correct document "
        "should print to the correct warehouse printer with no one clicking Print. This "
        "customisation delivers the reusable engine that renders a document and hands it to "
        "a printer once it is told to; deciding exactly when to fire, for which record, and "
        "at which printer remains the job of the NetSuite workflows built around it, so the "
        "same engine can support every scenario in the functional design without being "
        "rewritten for each one.",
    )

    doc.add_heading("1.4.2 Print Execution Engine", level=3)
    add_body(
        doc,
        "The Workflow Action Script is the component that actually produces a printed "
        "document. When a NetSuite workflow calls it after a record is saved, the script "
        "identifies the record that triggered it, works out which printed form to render "
        "for that record, and hands the resulting document off to the existing BB1 QZ Tray "
        "library so it can be delivered to a physical warehouse printer without anyone "
        "opening the record or clicking Print. Because the script does not contain any "
        "logic about when to fire, it can sit behind any of the workflows in the functional "
        "design's trigger matrix without being customer- or scenario-specific.",
    )

    doc.add_heading("1.4.3 Delivery via the Existing QZ Tray Library", level=3)
    add_body(
        doc,
        "Rather than talking to the warehouse printer directly, the action script queues "
        "the print request with BB1's existing QZ Tray library (already built and "
        "maintained separately) and redirects the user's browser back to the record. That "
        "library is responsible for picking the job up from its queue and delivering it to "
        "the physical printer through the QZ Tray application installed on the warehouse "
        "computer. This means the print engine did not need to reinvent the browser-to-"
        "printer bridge that the functional design's architecture called for - it reuses a "
        "mechanism that already exists.",
    )

    doc.add_heading("1.4.4 Configurability - Current State", level=3)
    add_body(
        doc,
        "The functional design asks for the engine to be configurable by document type, "
        "printer and quantity so it can be reused across every scenario in the trigger "
        "matrix, not just one. In its current, first-pass form the script only supports "
        "choosing which printed form to render, via a single script parameter set once on "
        "the deployment; it does not yet accept a separate document type, a specific target "
        "printer, or a print quantity per trigger. This is flagged as an open item in "
        "Section 4 rather than treated as finished, since closing that gap is what will let "
        "one script instance serve more than one of the workflows being built around it.",
    )

    add_page_break(doc)

    # ---------------- 2. Components ----------------
    doc.add_heading("2. Components", level=1)

    doc.add_heading("2.1 Custom Records", level=2)
    add_grey_note(
        doc,
        "None created as part of this deliverable. The Functional Design specifies a "
        "Printer custom record (fields _bb1_prn_name, _bb1_prn_ip_address) to hold "
        "warehouse printer details; it is not present in this SDF project and is not yet "
        "referenced by the script. See Section 2.7 for how it fits the wider design and "
        "Section 4, item 2, for its open ownership question.",
    )

    doc.add_heading("2.2 Custom Fields", level=2)
    add_grey_note(doc, "None created as part of this deliverable.")

    doc.add_heading("2.3 Custom Lists", level=2)
    add_grey_note(doc, "None created as part of this deliverable.")

    doc.add_heading("2.4 Scripts", level=2)
    add_grey_note(
        doc,
        "File cabinet path: src/FileCabinet/SuiteScripts/[JKR] Print Document/",
    )
    add_styled_table(
        doc,
        headers=["File Name", "Type", "Function", "Deployed To"],
        rows=[
            ["bb1_jkr_prnt_doc_wfa.js", "wfa",
             "Reads the triggering record's type and id from the workflow context, "
             "resolves a redirect URL back to that record, reads the configured print form "
             "number from a script parameter, builds a print job payload, and queues it "
             "with the existing BB1 QZ Tray library before redirecting the browser back to "
             "the record so the library can complete delivery.",
             "SuiteFlow Workflow Action step - workflow(s) configured separately by "
             "another developer (see Section 2.7)."],
        ],
        col_widths=[2400, 1000, 3600, 2638],
        mono_cols={0},
    )

    doc.add_heading("bb1_jkr_prnt_doc_wfa.js - Script Parameters", level=3)
    add_styled_table(
        doc,
        headers=["Field Name", "Field Internal ID", "Type", "Description"],
        rows=[
            ["Print Form Number", "custscript_bb1_jkr_auto_print_cust_form", "Integer/List",
             "Internal ID / number of the custom transaction form the script renders as "
             "the printed PDF. Currently a single value per script deployment - not yet a "
             "per-trigger Document Type, Printer or Quantity parameter as called for by "
             "the Functional Design (see Section 4, item 1)."],
        ],
        col_widths=[2500, 3000, 1600, 2538],
        mono_cols={1},
    )

    doc.add_heading("2.5 Other Files / External Dependencies", level=2)
    add_styled_table(
        doc,
        headers=["File Name", "Function"],
        rows=[
            ["/SuiteApps/com.bluebridgeonecouk.qztray/bb1_qz_lib_public",
             "Existing BB1 QZ Tray integration library (built and maintained separately "
             "by Wayne Wundram), loaded as a module dependency. Exposes AddJobToQueue() to "
             "queue a print job and InitiateProcessJobQueue() to process the queue and "
             "redirect the browser back to the source record, which is how the browser-to-"
             "QZ-Tray handoff described in the Functional Design's Open Points is actually "
             "resolved today. Its internals sit outside this SDF project."],
        ],
        col_widths=[3200, 6438],
        mono_cols={0},
    )

    doc.add_heading("2.6 Dependency Diagram", level=2)
    add_body(
        doc,
        "This component has no data model of its own - it is a single script and its "
        "dependencies. The diagram below shows the triggering record handing off to the "
        "action script, the script handing a print job to the existing BB1 QZ Tray "
        "library, and the Printer custom record from the Functional Design shown dashed, "
        "since it is specified but not yet built or referenced by this script.",
    )
    add_image(doc, "erd_wfa.png", width_in=6.3)

    doc.add_heading("2.7 Additional Components (Dependencies, Not Built Here)", level=2)
    doc.add_heading("SuiteFlow Workflows - built separately", level=3)
    add_body(
        doc,
        "Per the Functional Design's trigger matrix, the following NetSuite workflows are "
        "expected to call this action script once built. None of them are part of this "
        "deliverable; they are listed here so the script's parameters and behaviour can be "
        "reviewed against every scenario it needs to serve, not just the first one.",
    )
    add_styled_table(
        doc,
        headers=["Transaction & Trigger", "Trigger Type", "Condition Summary", "Output"],
        rows=[
            ["Sales Order / Released", "After Submit",
             "Order status is an open fulfilment state and at least one line has "
             "Quantity Committed > 0.", "Sales Order Pick Ticket"],
            ["Work Order-Build / Build saved", "After Submit",
             "Build is saved and the finished goods commit to an open Sales Order.",
             "Sales Order Pick Note"],
            ["Project Order / Created", "After Submit",
             "New Project Order created; a saved search prints the picking list for all "
             "its Work Orders.", "Picking Ticket"],
            ["Project Order / Event", "Custom Button",
             "Triggered per Gap 2 rules (creation or release of the Project Order).",
             "Manufacturing Picking Ticket / Manufacturing Document"],
            ["Item Fulfillment / Packed", "After Submit",
             "Fulfilment status changes to Picked and delivery method requires internal "
             "fleet labels.", "Address Labels (qty = Number of Boxes field)"],
        ],
        col_widths=[2400, 1400, 3800, 2038],
        mono_cols=set(),
    )

    doc.add_heading("Printer Custom Record - per Functional Design", level=3)
    add_body(
        doc,
        "The Functional Design specifies a custom record to hold each warehouse printer's "
        "name and IP address (fields _bb1_prn_name, _bb1_prn_ip_address). It is not present "
        "in this SDF project, and the current script does not read from it - the printer a "
        "job is delivered to is resolved entirely inside the existing BB1 QZ Tray library. "
        "Ownership of this record should be confirmed with whoever is building the "
        "workflows (see Section 4, item 2).",
    )

    doc.add_heading("2.8 Process Flow", level=2)
    add_body(
        doc,
        "The flow below traces what bb1_jkr_prnt_doc_wfa.js actually does today, from the "
        "point a (separately built) workflow calls it through to the existing QZ Tray "
        "library delivering the document to a printer.",
    )
    add_image(doc, "flow_wfa.png", width_in=6.3)

    add_page_break(doc)

    # ---------------- 3. Testing & Setup ----------------
    doc.add_heading("3. Testing & Setup", level=1)

    doc.add_heading("3.1 Steps to Test / Replicate", level=2)
    add_grey_note(
        doc,
        "Since the production trigger workflows are being built separately and were not "
        "available at the time of writing, testing this script in isolation requires a "
        "temporary test workflow - see step 2.",
    )
    add_numbered(doc, [
        "Confirm the BB1 QZ Tray SuiteApp (com.bluebridgeonecouk.qztray) is installed in "
        "the target account, and that QZ Tray is installed and running on the test machine "
        "with a printer registered and the browser's QZ Tray certificate/connection "
        "approved.",
        "Build a temporary SuiteFlow workflow on a test record type with an After Submit "
        "action calling bb1_jkr_prnt_doc_wfa.js, purely to exercise the script until the "
        "real trigger workflow(s) are delivered.",
        "Set the custscript_bb1_jkr_auto_print_cust_form parameter to a valid custom "
        "transaction form internal ID that exists on the target record type.",
        "Save/trigger the test record and confirm, from the script's existing log.debug "
        "statements, that Record Type, Record Id, Redirect Url and Form Number are all "
        "captured correctly.",
        "Confirm the browser redirects back to the record and the queued job - "
        "customscript_bb1_qz_pi_record, per the payload the script builds - is picked up "
        "by the BB1 QZ Tray library and the document reaches the configured printer.",
        "Repeat with an invalid or blank form number parameter and confirm the resulting "
        "behaviour; the script currently has no validation for this case (see Section 4, "
        "item 3).",
        "Remove the temporary test workflow once the real trigger workflow(s) are "
        "delivered, to avoid the same event printing twice.",
    ])

    doc.add_heading("3.2 Custom Record Setup", level=2)
    add_grey_note(
        doc,
        "Not applicable to this deliverable - no custom records were created. If the "
        "Printer custom record from the Functional Design (Section 2.7) is required, its "
        "setup steps belong to whoever builds it.",
    )

    add_page_break(doc)

    # ---------------- 4. Open Items / Gaps vs Functional Design ----------------
    doc.add_heading("4. Open Items / Gaps vs Functional Design", level=1)
    add_body(
        doc,
        "This is an initial pass at the technical notes, written against the script as it "
        "stands today rather than the full end state described in the Functional Design. "
        "The items below are the known gaps and open questions, so they are visible before "
        "the dependent workflows are built on top of this script.",
    )
    add_styled_table(
        doc,
        headers=["#", "Item", "Category", "Notes"],
        rows=[
            ["1", "Document Type, Printer and Quantity are not yet separate parameters",
             "Gap vs Design",
             "The Functional Design calls for the engine to be configured by Document "
             "Type, Printer and Quantity per trigger. Today there is only one script "
             "parameter (print form number), fixed per deployment, so one script "
             "instance cannot yet serve more than one document type/printer/quantity "
             "combination."],
            ["2", "Printer custom record ownership is unconfirmed", "Decision Needed",
             "Specified in the Functional Design (_bb1_prn_name, _bb1_prn_ip_address) but "
             "not present in this SDF project and not referenced by the script. Confirm "
             "with the developer building the workflows whether it is needed and who "
             "owns it."],
            ["3", "No validation on the form number parameter", "Gap vs Design",
             "onAction() reads custscript_bb1_jkr_auto_print_cust_form and uses it "
             "directly with no check for missing or invalid values before queuing the "
             "print job."],
            ["4", "Script does not branch by document type", "Gap vs Design",
             "newRecord.type/id are read generically, but the script does not vary its "
             "behaviour by which of the five trigger-matrix scenarios called it, so every "
             "trigger currently renders the same hardcoded form."],
            ["5", "Quantity (e.g. Number of Boxes for Address Labels) is not implemented",
             "Gap vs Design", "jobData has no explicit quantity field; the Address "
             "Labels scenario in the trigger matrix needs one copy per box."],
            ["6", "BB1 QZ Library internals are outside this project", "Dependency",
             "AddJobToQueue()/InitiateProcessJobQueue() are assumed to resolve the "
             "Functional Design's Open Point #1 (server-to-client handoff), but the "
             "library's source is not in this SDF project - confirm its exact behaviour "
             "with Wayne Wundram before relying on it for every trigger scenario."],
            ["7", "Script has not been exercised against all five trigger-matrix rows",
             "Test Coverage",
             "Only tested (per Section 3.1) via a temporary workflow; needs to be "
             "re-verified once the real workflows - including the Work Order/Build "
             "scenario - are built and wired up."],
        ],
        col_widths=[500, 2900, 1500, 4738],
        mono_cols=set(),
    )

    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()
